"""Exact image anchors; source DOCX formatting is retained."""
from __future__ import annotations
import hashlib
import re
import zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

MARKER = re.compile(r"\[\[IMAGE:([A-Za-z][A-Za-z0-9_-]{0,63})\]\]")

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def read_manuscript(path):
    path = Path(path)
    if path.stat().st_size > 2_000_000:
        raise ValueError('MANUSCRIPT_TOO_LARGE')
    if path.suffix == '.docx':
        with zipfile.ZipFile(path) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40_000_000:
                raise ValueError('MANUSCRIPT_ARCHIVE_TOO_LARGE')
        doc = Document(path)
        # Table/header/text-box anchors cannot be replaced safely by body indices.
        direct = {p._p for p in doc.paragraphs}
        for node in doc.element.body.xpath('.//w:p'):
            text = ''.join(node.xpath('.//w:t/text()'))
            if '[[IMAGE:' in text and node not in direct:
                raise ValueError('MARKER_MUST_BE_IN_BODY_PARAGRAPH')
        blocks = [''.join(p._p.xpath('.//w:t/text()')) for p in doc.paragraphs]
    elif path.suffix in {'.md', '.txt'}:
        blocks, pending = [], []
        for line in path.read_text(encoding='utf-8-sig').splitlines():
            if MARKER.fullmatch(line.strip()) or not line.strip():
                if pending:
                    blocks.append('\n'.join(pending)); pending = []
                if line.strip(): blocks.append(line.strip())
            else:
                pending.append(line)
        if pending: blocks.append('\n'.join(pending))
    else:
        raise ValueError('MANUSCRIPT_FORMAT_INVALID')
    anchors, seen, prefix = [], set(), []
    for i, text in enumerate(blocks):
        match = MARKER.fullmatch(text.strip())
        if '[[IMAGE:' in text and not match:
            raise ValueError('MARKER_MUST_OCCUPY_ITS_OWN_PARAGRAPH')
        if match:
            scene_id = match[1]
            if scene_id in seen: raise ValueError('DUPLICATE_MARKER')
            if not any(p.strip() for p in prefix): raise ValueError('MARKER_BEFORE_STORY')
            seen.add(scene_id)
            anchors.append({'scene_id':scene_id, 'paragraph_index':i,
                            'context':'\n\n'.join(prefix)[-7000:]})
        else:
            prefix.append(text)
    if not anchors: raise ValueError('NO_IMAGE_MARKERS')
    if len(anchors) > 12: raise ValueError('TOO_MANY_IMAGE_MARKERS')
    return blocks, anchors

def assemble_marked(source, insertions, output):
    blocks, anchors = read_manuscript(source)
    if set(insertions) != {a['scene_id'] for a in anchors}:
        raise ValueError('IMAGE_MARKER_SET_MISMATCH')
    is_docx = Path(source).suffix == '.docx'
    doc = Document(source) if is_docx else Document()
    if not is_docx:
        doc.styles['Normal'].font.name = 'Georgia'
        doc.styles['Normal'].font.size = Pt(11)
        doc.styles['Normal'].paragraph_format.space_after = Pt(8)
        for name,size in [('Title',24),('Heading 1',16)]:
            doc.styles[name].font.name='Georgia'
            doc.styles[name].font.size=Pt(size)
            doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        for border in doc.styles.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        for text in blocks:
            if text.startswith('# '): doc.add_paragraph(text[2:], 'Title')
            elif text.startswith('## '): doc.add_paragraph(text[3:], 'Heading 1')
            else: doc.add_paragraph(text)
    for anchor in anchors:
        paragraph = doc.paragraphs[anchor['paragraph_index']]
        paragraph.clear()
        paragraph.alignment = 1
        paragraph.paragraph_format.keep_together = True
        # Respect the section containing this anchor, including landscape sections.
        section = doc.sections[-1]
        next_section = paragraph._p.xpath('following::w:sectPr')
        if next_section:
            from docx.section import Section
            section = Section(next_section[0], doc.part)
        width = (section.page_width-section.left_margin-section.right_margin)/914400
        height = min(4.8,(section.page_height-section.top_margin-section.bottom_margin)/914400*.55)
        image_path = insertions[anchor['scene_id']]
        with Image.open(image_path) as img: iw, ih = img.size
        inches = min(width, height*iw/ih)
        run = paragraph.add_run()
        picture = run.add_picture(str(image_path), width=Inches(inches))
        picture._inline.docPr.set('descr', 'Illustration '+anchor['scene_id'])
    out = Path(output); out.parent.mkdir(parents=True, exist_ok=True)
    temporary = out.with_suffix('.partial.docx')
    doc.save(temporary); temporary.replace(out)

