import base64
import io
import json
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
from pathlib import Path
from PIL import Image
from docx import Document
from content_profile import validate_profile
from marked_engine import MarkedEngine, JobError
from marked_manuscript import read_manuscript, assemble_marked
from workbench import make_server

PROFILE={'version':1,'content_type':'mystery','format':'portrait','characters':{},'locations':{},'scenes':{}}
STORY='An adult archivist enters a quiet room.\n\n[[IMAGE:first]]\n\nShe finds a letter.\n\n[[IMAGE:second]]'
OWNER='a'*64
class FakeEdge:
    def __init__(self): self.descriptions=0; self.generations=0; self.fail=False
    def describe(self,anchor,profile):
        self.descriptions+=1
        return 'A fully clothed adult archivist stands by a quiet archive table.',{'provider':'fixture','model':'fixture'}
    def generate(self,prompt,negative,path,width,height):
        self.generations+=1
        if self.fail: raise RuntimeError('PROVIDER_RESPONSE_UNCERTAIN')
        Image.new('RGB',(width,height),(116,133,145)).save(path)
        return {'model':'fixture'}

class MarkedSystemTest(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.edge=FakeEdge();self.engine=MarkedEngine(self.root/'jobs',edge=self.edge)
    def tearDown(self): self.temp.cleanup()
    def create(self,story=STORY): return self.engine.create(OWNER,'story.md',story.encode(),PROFILE)
    def test_exact_markers_and_no_future_leak(self):
        job=self.create();stored=self.engine.get(job['id'],OWNER)
        self.assertNotIn('finds a letter',stored['images'][0]['context'])
        self.assertIn('finds a letter',stored['images'][1]['context'])
        for text in ['[[IMAGE:a]]\n\nText','Text\n\n[[IMAGE:a]]\n\n[[IMAGE:a]]','Text [[IMAGE:a]]']:
            with self.assertRaises(ValueError): self.create(text)
    def test_document_preserves_order_and_approval(self):
        job=self.create();jid=job['id']
        with self.assertRaisesRegex(JobError,'APPROVAL'):self.engine.assemble(jid,OWNER)
        result=self.engine.generate(jid,OWNER)
        self.assertEqual([x['status'] for x in result['images']],['review_required']*2)
        self.engine.generate(jid,OWNER);self.assertEqual(self.edge.generations,2)
        for scene in ['first','second']: self.engine.approve(jid,OWNER,scene)
        self.engine.assemble(jid,OWNER)
        doc=Document(self.root/'jobs'/jid/'illustrated.docx')
        self.assertEqual(len(doc.inline_shapes),2)
        self.assertNotIn('[[IMAGE:','\n'.join(p.text for p in doc.paragraphs))
        body=doc.element.body.xml
        self.assertLess(body.index('enters a quiet room'),body.index('w:drawing'))
        self.assertLess(body.index('w:drawing'),body.index('finds a letter'))
    def test_docx_format_and_table_anchor(self):
        source=self.root/'source.docx';doc=Document();p=doc.add_paragraph();p.add_run('Bold source').bold=True
        doc.add_paragraph('[[IMAGE:frame]]');doc.add_paragraph('After the picture');doc.save(source)
        image=self.root/'frame.png';Image.new('RGB',(768,1024),'white').save(image)
        out=self.root/'out.docx';assemble_marked(source,{'frame':image},out)
        saved=Document(out);self.assertTrue(saved.paragraphs[0].runs[0].bold)
        self.assertEqual(saved.paragraphs[2].text,'After the picture')
        doc.add_table(rows=1,cols=1).cell(0,0).text='[[IMAGE:table]]';doc.save(source)
        with self.assertRaisesRegex(ValueError,'BODY_PARAGRAPH'):read_manuscript(source)
    def test_uncertain_requests_not_replayed(self):
        job=self.create();jid=job['id'];self.edge.fail=True
        self.engine.generate(jid,OWNER);self.assertEqual(self.edge.generations,1)
        self.engine.generate(jid,OWNER);self.assertEqual(self.edge.generations,2,'Other not-yet-submitted scene may be submitted once')
        self.engine.generate(jid,OWNER);self.assertEqual(self.edge.generations,2)
        self.edge.fail=False;self.engine.generate(jid,OWNER,retry=True)
        self.assertEqual(self.edge.descriptions,2,'Saved prompts are reused')
        self.assertEqual(self.edge.generations,4)
    def test_restart_and_ownership(self):
        job=self.create();jid=job['id'];stored=self.engine.get(jid,OWNER)
        stored['images'][0]['status']='generating';self.engine._save(stored)
        restarted=MarkedEngine(self.root/'jobs',edge=self.edge)
        self.assertEqual(restarted.get(jid,OWNER)['images'][0]['status'],'uncertain')
        with self.assertRaisesRegex(JobError,'NOT_FOUND'):restarted.get(jid,'b'*64)
        self.assertEqual(restarted.list('b'*64),[])
    def test_mutated_images_and_sources_block_export(self):
        jid=self.create()['id'];self.engine.generate(jid,OWNER)
        for scene in ['first','second']:self.engine.approve(jid,OWNER,scene)
        Image.new('RGB',(768,1024),'red').save(self.root/'jobs'/jid/'first.png')
        with self.assertRaisesRegex(JobError,'IMAGE_CHANGED'):self.engine.assemble(jid,OWNER)
        (self.root/'jobs'/jid/'source.md').write_text('Changed')
        with self.assertRaisesRegex(JobError,'SOURCE_CHANGED'):self.engine.assemble(jid,OWNER)
    def test_refusal_not_replayed(self):
        def reject(*args): raise RuntimeError('CONTENT_NOT_SUPPORTED')
        self.edge.describe=reject;jid=self.create()['id']
        self.engine.generate(jid,OWNER);self.engine.generate(jid,OWNER,retry=True)
        self.assertTrue(all(x['status']=='blocked' for x in self.engine.get(jid,OWNER)['images']))
        self.assertEqual(self.edge.generations,0)
    def test_profile_cannot_set_backend_or_execute(self):
        with self.assertRaisesRegex(ValueError,'UNKNOWN_FIELD'): validate_profile({**PROFILE,'backend_url':'http://private'})
        with self.assertRaisesRegex(ValueError,'TYPE_INVALID'):validate_profile({**PROFILE,'content_type':'unrestricted'})
    def test_parallel_generation_excluded(self):
        entered,release=threading.Event(),threading.Event();base=self.edge.generate
        def delayed(*args):entered.set();release.wait(3);return base(*args)
        self.edge.generate=delayed;jid=self.create()['id']
        thread=threading.Thread(target=lambda:self.engine.generate(jid,OWNER));thread.start();entered.wait(2)
        try:
            with self.assertRaisesRegex(JobError,'BUSY'):self.engine.generate(jid,OWNER)
            self.assertEqual(self.engine.get(jid,OWNER)['images'][0]['status'],'generating')
        finally:release.set();thread.join(5)
    def test_service_authentication(self):
        server=make_server(self.engine,'t'*64,port=0);thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        url='http://127.0.0.1:'+str(server.server_port)+'/list'
        try:
            request=urllib.request.Request(url,data=b'{}',method='POST')
            with self.assertRaises(urllib.error.HTTPError) as error:urllib.request.urlopen(request)
            self.assertEqual(error.exception.code,403)
            request=urllib.request.Request(url,data=b'{}',headers={'Authorization':'Bearer '+'t'*64,'X-Workbench-Owner':OWNER},method='POST')
            with urllib.request.urlopen(request) as response:self.assertEqual(json.load(response),{'ok':True,'jobs':[]})
        finally:server.shutdown();server.server_close()

if __name__=='__main__':unittest.main()
