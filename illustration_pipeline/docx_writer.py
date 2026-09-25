from __future__ import annotations
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement

def read_docx_paragraphs(path: str | Path) -> list[str]:
    doc = Document(str(path))
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]

def build_docx(paragraphs: list[str], insertions: dict[int, str], output: str | Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    usable_inches = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
    image_width = Inches(min(float(usable_inches), 6.35))

    for idx, text in enumerate(paragraphs):
        p = doc.add_paragraph(text)
        if idx in insertions:
            pic = doc.add_paragraph()
            pic.alignment = 1
            pic.add_run().add_picture(insertions[idx], width=image_width)
    out = Path(output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))

def insert_into_existing_docx(source: str | Path, insertions: dict[int, str], output: str | Path) -> None:
    doc = Document(str(source))
    visible = [p for p in doc.paragraphs if p.text.strip()]
    sec = doc.sections[0]
    usable_inches = (sec.page_width - sec.left_margin - sec.right_margin) / 914400
    image_width = Inches(min(float(usable_inches), 6.35))

    for idx in sorted(insertions, reverse=True):
        if idx < 0 or idx >= len(visible):
            continue
        anchor = visible[idx]
        temp = doc.add_paragraph()
        temp.alignment = 1
        temp.add_run().add_picture(insertions[idx], width=image_width)
        anchor._p.addnext(deepcopy(temp._p))
        temp._element.getparent().remove(temp._element)

    out = Path(output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
